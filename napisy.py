import base64
import html
import io
import json
import math
import os
import re
import urllib.parse
import urllib.request
from openai import OpenAI
import streamlit as st
import streamlit.components.v1 as components
from youtube_transcript_api import YouTubeTranscriptApi

# Opcjonalne biblioteki do eksportu DOCX i PDF
try:
  from docx import Document

  DOCX_AVAILABLE = True
except ImportError:
  DOCX_AVAILABLE = False

try:
  from fpdf import FPDF

  FPDF_AVAILABLE = True
except ImportError:
  FPDF_AVAILABLE = False

APP_VERSION = "3.2"

# Bezpieczne pobieranie klucza z secrets
DEEPSEEK_API_KEY = st.secrets.get("DEEPSEEK_API_KEY", "")
DEEPSEEK_MODEL = "deepseek-chat"

st.set_page_config(
    page_title=f"YouTube Transkrypcja & DeepSeek AI v{APP_VERSION}",
    page_icon="🎬",
    layout="wide",
)


def render_sidebar_logo():
  """Wczytuje i wyświetla logo3.svg jako klikalny baner prowadzący do fabryka.tech."""
  svg_filename = "logo3.svg"
  script_dir = os.path.dirname(os.path.abspath(__file__))
  potential_paths = [
      svg_filename,
      os.path.join(script_dir, svg_filename),
  ]

  logo_path = None
  for p in potential_paths:
    if os.path.exists(p):
      logo_path = p
      break

  if logo_path:
    with open(logo_path, "rb") as f:
      svg_base64 = base64.b64encode(f.read()).decode("utf-8")

    st.sidebar.markdown(
        f"""
            <div style="text-align: center; margin-bottom: 20px;">
                <a href="https://fabryka.tech" target="_blank" rel="noopener noreferrer">
                    <img src="data:image/svg+xml;base64,{svg_base64}" alt="Fabryka Tech" style="max-width: 100%; height: auto; max-height: 70px; display: block; margin: 0 auto; filter: drop-shadow(0px 2px 4px rgba(0,0,0,0.2));" />
                </a>
            </div>
            """,
        unsafe_allow_html=True,
    )
  else:
    st.sidebar.markdown(
        """
            <div style="text-align: center; margin-bottom: 20px;">
                <a href="https://fabryka.tech" target="_blank" style="font-size: 1.3rem; font-weight: bold; text-decoration: none; color: #ff4b4b;">
                    🚀 FABRYKA.TECH
                </a>
            </div>
            """,
        unsafe_allow_html=True,
    )


# Inicjalizacja stanu sesji
if "history" not in st.session_state:
  st.session_state["history"] = []
if "transcript_text" not in st.session_state:
  st.session_state["transcript_text"] = ""
if "video_id" not in st.session_state:
  st.session_state["video_id"] = ""
if "video_title" not in st.session_state:
  st.session_state["video_title"] = ""


# --- PASEK BOCZNY (SIDEBAR) ---
render_sidebar_logo()

st.sidebar.markdown("---")
st.sidebar.header("🕒 Historia sesji (max 5)")
if not st.session_state["history"]:
  st.sidebar.info("Brak pobranych filmów w bieżącej sesji.")
else:
  for idx, item in enumerate(st.session_state["history"]):
    label = f"🎬 {item['title'][:26]}..."
    if st.sidebar.button(
        label,
        key=f"hist_{idx}_{item['video_id']}",
        use_container_width=True,
    ):
      st.session_state["transcript_text"] = item["text"]
      st.session_state["video_id"] = item["video_id"]
      st.session_state["video_title"] = item["title"]
      st.rerun()

st.sidebar.markdown("---")
st.sidebar.link_button(
    "📩 Formularz kontaktowy",
    "https://fabryka.tech/kontakt",
    use_container_width=True,
)


# --- GŁÓWNA STRONA ---
st.title(f"🎬 YouTube Transkrypcja do tekstu & DeepSeek AI (v{APP_VERSION})")
st.write(
    "Pobieraj transkrypcje, przeglądaj podgląd wideo, eksportuj pliki i generuj"
    " inteligentne analizy z DeepSeek AI."
)

with st.expander("ℹ️ Dlaczego to narzędzie powstało?"):
  st.markdown("""
    Różne modele językowe, jak Gemini czy ChatGPT, mają problem z pobraniem transkrypcji do filmów z powodów:
    * **Blokad antybotowych i limitów żądań (Error 429)** nakładanych przez YouTube na adresy IP serwerów AI,
    * **Braku obsługi sesji i dynamicznego odtwarzacza**, przez co boty nie odczytują automatycznie generowanych napisów,
    * **Zmiennych formatów napisów (JSON3/XML)**, których standardowe parsery stron internetowych nie przetwarzają.
    
    Dlatego przychodzimy z pomocą w postaci łatwego pobierania napisów – wyciągamy tekst z opcją formatowania akapitów, statystykami, historią, odtwarzaczem i bezpośrednią analizą DeepSeek AI.
    """)


def sanitize_filename(title: str, max_len: int = 50) -> str:
  clean = re.sub(r'[\\/*?:"<>|]', "", title)
  clean = clean.replace(" ", "_").strip("._ ")
  return clean[:max_len] if clean else "film"


def get_video_title(video_id: str) -> str:
  try:
    url = f"https://www.youtube.com/watch?v={video_id}"
    oembed_url = f"https://www.youtube.com/oembed?url={urllib.parse.quote(url)}&format=json"
    req = urllib.request.Request(
        oembed_url, headers={"User-Agent": "Mozilla/5.0"}
    )
    with urllib.request.urlopen(req, timeout=4) as response:
      data = json.loads(response.read().decode("utf-8"))
      return data.get("title", f"video_{video_id}")
  except Exception:
    return f"video_{video_id}"


def add_to_history(video_id: str, title: str, text: str):
  st.session_state["history"] = [
      item
      for item in st.session_state["history"]
      if item["video_id"] != video_id
  ]
  st.session_state["history"].insert(
      0, {"video_id": video_id, "title": title, "text": text}
  )
  st.session_state["history"] = st.session_state["history"][:5]


def format_seconds(seconds: float) -> str:
  total_seconds = int(seconds)
  hours = total_seconds // 3600
  minutes = (total_seconds % 3600) // 60
  secs = total_seconds % 60
  if hours > 0:
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"
  return f"{minutes:02d}:{secs:02d}"


def extract_video_id(url: str) -> str:
  patterns = [
      r"(?:v=|\/)([0-9A-Za-z_-]{11}).*",
      r"(?:youtu\.be\/)([0-9A-Za-z_-]{11})",
  ]
  for pattern in patterns:
    match = re.search(pattern, url)
    if match:
      return match.group(1)
  return ""


def fetch_raw_transcript(video_id: str, preferred_lang: str = "pl") -> list:
  api = YouTubeTranscriptApi()
  try:
    transcript_list = api.list(video_id)
  except AttributeError:
    transcript_list = YouTubeTranscriptApi.list_transcripts(video_id)

  target_transcript = None
  for t in transcript_list:
    if t.language_code.lower().startswith(preferred_lang.lower()):
      target_transcript = t
      break

  if not target_transcript:
    for t in transcript_list:
      if t.is_translatable:
        target_transcript = t.translate(preferred_lang)
        break
      else:
        target_transcript = t
        break

  if not target_transcript:
    raise ValueError("Nie znaleziono żadnych napisów dla tego materiału.")

  return target_transcript.fetch()


def count_sentences(text: str) -> int:
  return len(re.findall(r"[.!?]+(?:\s|$)", text))


def split_into_paragraphs(
    entries: list, pause_threshold: float = 1.5, min_sentences: int = 5
) -> list[str]:
  paragraphs = []
  current_sentences = []
  current_sentence_count = 0

  for i, entry in enumerate(entries):
    txt = entry["text"]
    current_sentences.append(txt)
    current_sentence_count += count_sentences(txt)

    is_long_pause = False
    if i < len(entries) - 1:
      next_entry = entries[i + 1]
      prev_end = entry["start"] + entry.get("duration", 0.0)
      gap = next_entry["start"] - prev_end
      if gap >= pause_threshold:
        is_long_pause = True

    if (
        is_long_pause and current_sentence_count >= min_sentences
    ) or current_sentence_count >= min_sentences * 2:
      paragraphs.append(" ".join(current_sentences))
      current_sentences = []
      current_sentence_count = 0

  if current_sentences:
    paragraphs.append(" ".join(current_sentences))

  return paragraphs


def format_transcript(
    data: list,
    include_time: bool,
    include_numbering: bool,
    use_smart_paragraphs: bool,
) -> str:
  entries = []
  seen_texts = set()

  for item in data:
    text = (
        item.get("text", "")
        if isinstance(item, dict)
        else getattr(item, "text", "")
    )
    start = (
        item.get("start", 0.0)
        if isinstance(item, dict)
        else getattr(item, "start", 0.0)
    )
    duration = (
        item.get("duration", 0.0)
        if isinstance(item, dict)
        else getattr(item, "duration", 0.0)
    )
    text = text.strip()

    if not text or text in seen_texts:
      continue
    seen_texts.add(text)
    entries.append({"text": text, "start": start, "duration": duration})

  if use_smart_paragraphs and not include_time and not include_numbering:
    paragraphs = split_into_paragraphs(
        entries, pause_threshold=1.5, min_sentences=5
    )
    return "\n\n".join(paragraphs)

  if not include_time and not include_numbering:
    return " ".join([entry["text"] for entry in entries])

  lines = []
  for idx, entry in enumerate(entries, start=1):
    prefix_parts = []
    if include_numbering:
      prefix_parts.append(f"{idx}.")
    if include_time:
      prefix_parts.append(f"[{format_seconds(entry['start'])}]")

    prefix = " ".join(prefix_parts)
    lines.append(f"{prefix} {entry['text']}".strip())

  return "\n".join(lines)


def create_docx_bytes(text: str, title: str) -> bytes:
  doc = Document()
  doc.add_heading(f"Transkrypcja: {title}", level=1)
  for p in text.split("\n\n"):
    if p.strip():
      doc.add_paragraph(p.strip())
  doc_io = io.BytesIO()
  doc.save(doc_io)
  return doc_io.getvalue()


def create_pdf_bytes(text: str, title: str) -> bytes:
  pdf = FPDF()
  pdf.add_page()
  pdf.set_auto_page_break(auto=True, margin=15)

  font_path = "C:/Windows/Fonts/arial.ttf"
  if os.path.exists(font_path):
    pdf.add_font("ArialUnicode", "", font_path)
    pdf.set_font("ArialUnicode", size=11)
    font_name = "ArialUnicode"
  else:
    pdf.set_font("helvetica", size=11)
    font_name = "helvetica"

  pdf.set_font(font_name, size=13)
  pdf.cell(
      0,
      10,
      f"Transkrypcja: {title[:70]}",
      new_x="LMARGIN",
      new_y="NEXT",
  )
  pdf.ln(4)

  pdf.set_font(font_name, size=10)
  for paragraph in text.split("\n\n"):
    if paragraph.strip():
      pdf.multi_cell(0, 6, paragraph.strip())
      pdf.ln(3)

  return bytes(pdf.output())


def render_clipboard_button(text_to_copy: str):
  clean_js_text = json.dumps(text_to_copy)
  copy_html = f"""
    <div style="margin-top: 5px; margin-bottom: 10px;">
        <button id="copy-btn" onclick="copyAction()" style="
            background-color: #262730;
            color: #ffffff;
            border: 1px solid #4a4d5a;
            border-radius: 8px;
            padding: 8px 16px;
            font-size: 14px;
            font-weight: 500;
            cursor: pointer;
            transition: all 0.2s ease;
            display: inline-flex;
            align-items: center;
            gap: 6px;
        ">📋 Kopiuj do schowka</button>
    </div>
    <script>
    function copyAction() {{
        const text = {clean_js_text};
        const btn = document.getElementById('copy-btn');
        if (navigator.clipboard && window.isSecureContext) {{
            navigator.clipboard.writeText(text).then(() => showSuccess(btn));
        }} else {{
            const ta = document.createElement('textarea');
            ta.value = text;
            document.body.appendChild(ta);
            ta.select();
            document.execCommand('copy');
            document.body.removeChild(ta);
            showSuccess(btn);
        }}
    }}
    function showSuccess(btn) {{
        btn.innerText = '✅ Skopiowano!';
        btn.style.backgroundColor = '#1e7e34';
        btn.style.borderColor = '#28a745';
        setTimeout(() => {{
            btn.innerText = '📋 Kopiuj do schowka';
            btn.style.backgroundColor = '#262730';
            btn.style.borderColor = '#4a4d5a';
        }}, 2500);
    }}
    </script>
    """
  components.html(copy_html, height=50)


def summarize_with_deepseek(text: str) -> str:
  if not DEEPSEEK_API_KEY:
    raise ValueError(
        "Brak klucza API w st.secrets. Dodaj DEEPSEEK_API_KEY w ustawieniach"
        " secrets aplikacji."
    )

  client = OpenAI(
      api_key=DEEPSEEK_API_KEY, base_url="https://api.deepseek.com"
  )
  prompt = f"""Przeanalizuj poniższą transkrypcję z filmu YouTube i przygotuj wyczerpujące, czytelne podsumowanie w języku polskim.
Struktura podsumowania:
1. Główne przesłanie (krótki wstęp / sedno sprawy)
2. Kluczowe punkty i wnioski (w punktach)
3. Najważniejsza konkluzja / puenta

Treść transkrypcji:
{text[:45000]}
"""
  response = client.chat.completions.create(
      model=DEEPSEEK_MODEL,
      messages=[
          {
              "role": "system",
              "content": (
                  "Jesteś precyzyjnym asystentem analitycznym tworzącym wysokiej"
                  " jakości streszczenia materiałów wideo w języku polskim."
              ),
          },
          {"role": "user", "content": prompt},
      ],
      stream=False,
  )
  return response.choices[0].message.content


# FORMULARZ POBIERANIA
url_input = st.text_input(
    "Wklej link do filmu YouTube:",
    placeholder="https://www.youtube.com/watch?v=...",
)

col1, col2, col3, col4 = st.columns([1, 1, 1, 1.2])
with col1:
  selected_lang = st.selectbox(
      "Preferowany język:",
      options=["pl", "en", "de", "es", "fr"],
      index=0,
  )
with col2:
  include_time = st.checkbox("Dodaj znaczniki czasu", value=False)
with col3:
  include_numbering = st.checkbox("Dodaj numerację linii", value=False)
with col4:
  use_smart_paragraphs = st.checkbox(
      "Dziel na akapity (min. 5 zdań + pauzy)",
      value=True,
      help=(
          "Tworzy nowy akapit przy przerwach w mówieniu, zachowując min. 5 zdań"
          " na akapit."
      ),
  )

if st.button("Pobierz transkrypcję", type="primary"):
  video_id = extract_video_id(url_input.strip())
  if not video_id:
    st.warning("Podaj poprawny link do filmu YouTube.")
  else:
    with st.spinner("Pobieranie tytułu i transkrypcji..."):
      try:
        title = get_video_title(video_id)
        raw_data = fetch_raw_transcript(video_id, selected_lang)
        formatted_text = format_transcript(
            raw_data,
            include_time,
            include_numbering,
            use_smart_paragraphs,
        )

        st.session_state["transcript_text"] = formatted_text
        st.session_state["video_id"] = video_id
        st.session_state["video_title"] = title

        add_to_history(video_id, title, formatted_text)
        st.success(f"Transkrypcja pobrana pomyślnie: **{title}**")
      except Exception as e:
        st.error(f"Wystąpił błąd: {e}")

# WYŚWIETLANIE WYNIKÓW
if st.session_state["transcript_text"]:
  text_data = st.session_state["transcript_text"]
  v_id = st.session_state["video_id"]
  v_title = st.session_state.get("video_title", f"film_{v_id}")
  file_base = sanitize_filename(v_title)

  st.info(f"Aktualnie załadowany materiał: **{v_title}**")

  # --- 1. STATYSTYKI TEKSTU ---
  st.divider()
  words = text_data.split()
  word_count = len(words)
  char_count = len(text_data)
  char_count_no_spaces = len(text_data.replace(" ", "").replace("\n", ""))
  reading_time = max(1, math.ceil(word_count / 200))

  st.subheader("📊 Statystyki transkrypcji")
  m1, m2, m3, m4 = st.columns(4)
  m1.metric("Liczba słów", f"{word_count:,}".replace(",", " "))
  m2.metric("Znaki (ze spacjami)", f"{char_count:,}".replace(",", " "))
  m3.metric("Znaki (bez spacji)", f"{char_count_no_spaces:,}".replace(",", " "))
  m4.metric("Szacowany czas czytania", f"~{reading_time} min")

  st.divider()

  # --- 2. GŁÓWNA SEKCJA: WIDEO + EDYTOR + FORMATY + DEEPSEEK AI ---
  col_left, col_right = st.columns(2)

  with col_left:
    st.subheader("🎬 Podgląd materiału:")
    video_url = f"https://www.youtube.com/watch?v={v_id}"
    st.video(video_url)

    st.subheader("📄 Pobrana transkrypcja:")
    st.text_area(
        label="Treść",
        value=text_data,
        height=350,
        label_visibility="collapsed",
    )

    render_clipboard_button(text_data)

    st.write("💾 **Pobierz jako plik o tytule filmu:**")
    d_col1, d_col2, d_col3, d_col4 = st.columns(4)

    with d_col1:
      st.download_button(
          label="📄 TXT",
          data=text_data,
          file_name=f"{file_base}.txt",
          mime="text/plain",
          use_container_width=True,
      )

    with d_col2:
      st.download_button(
          label="📝 Markdown",
          data=f"# {v_title}\n\n{text_data}",
          file_name=f"{file_base}.md",
          mime="text/markdown",
          use_container_width=True,
      )

    with d_col3:
      if DOCX_AVAILABLE:
        docx_file = create_docx_bytes(text_data, v_title)
        st.download_button(
            label="📘 Word (.docx)",
            data=docx_file,
            file_name=f"{file_base}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )
      else:
        st.button(
            "📘 DOCX (brak)",
            disabled=True,
            help="Zainstaluj python-docx: pip install python-docx",
            use_container_width=True,
        )

    with d_col4:
      if FPDF_AVAILABLE:
        pdf_file = create_pdf_bytes(text_data, v_title)
        st.download_button(
            label="📕 PDF",
            data=pdf_file,
            file_name=f"{file_base}.pdf",
            mime="application/pdf",
            use_container_width=True,
        )
      else:
        st.button(
            "📕 PDF (brak)",
            disabled=True,
            help="Zainstaluj fpdf2: pip install fpdf2",
            use_container_width=True,
        )

    # --- 3. WYSZUKIWARKA W TEKŚCIE ---
    st.markdown("---")
    st.subheader("🔍 Wyszukiwarka w tekście")
    search_query = st.text_input(
        "Wpisz szukane słowo lub frazę:",
        placeholder="np. sztuczna inteligencja",
        key="search_input",
    )

    if search_query.strip():
      query_clean = search_query.strip()
      occurrences = len(
          re.findall(re.escape(query_clean), text_data, flags=re.IGNORECASE)
      )

      if occurrences > 0:
        st.success(
            f"Znaleziono **{occurrences}** wystąpień frazy: *„{query_clean}”*"
        )

        def highlight_match(match):
          return (
              "<mark style='background-color: #ffe066; color: #000; padding:"
              f" 2px 4px; border-radius: 3px;'>{match.group(0)}</mark>"
          )

        highlighted_text = re.sub(
            re.escape(query_clean),
            highlight_match,
            html.escape(text_data),
            flags=re.IGNORECASE,
        )
        highlighted_html = highlighted_text.replace("\n\n", "<br><br>").replace(
            "\n", "<br>"
        )

        st.markdown(
            f"""<div style='max-height: 250px; overflow-y: auto; padding: 12px; border: 1px solid #444; border-radius: 8px; font-size: 0.95rem; line-height: 1.6; background-color: rgba(255,255,255,0.03);'>
                    {highlighted_html}
                    </div>""",
            unsafe_allow_html=True,
        )
      else:
        st.warning(f"Nie znaleziono słowa: *„{query_clean}”* w tekście.")

  with col_right:
    st.subheader("🤖 Analiza DeepSeek AI")
    if st.button("✨ Generuj podsumowanie AI", type="secondary"):
      if not DEEPSEEK_API_KEY:
        st.error(
            "Nie skonfigurowano klucza DeepSeek API w Streamlit Secrets!"
            " Uzupełnij pole `DEEPSEEK_API_KEY` w panelu hostingu lub w"
            " `.streamlit/secrets.toml`."
        )
      else:
        with st.spinner("Przetwarzanie przez DeepSeek AI..."):
          try:
            summary = summarize_with_deepseek(text_data)
            st.markdown(summary)

            s_col1, s_col2 = st.columns(2)
            with s_col1:
              st.download_button(
                  label="💾 Podsumowanie (.md)",
                  data=summary,
                  file_name=f"{file_base}_podsumowanie.md",
                  mime="text/markdown",
                  use_container_width=True,
              )
            with s_col2:
              st.download_button(
                  label="💾 Podsumowanie (.txt)",
                  data=summary,
                  file_name=f"{file_base}_podsumowanie.txt",
                  mime="text/plain",
                  use_container_width=True,
              )
          except Exception as e:
            st.error(f"Błąd DeepSeek API: {e}")

st.caption(f"Wersja aplikacji: v{APP_VERSION}")